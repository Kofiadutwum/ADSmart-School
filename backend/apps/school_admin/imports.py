"""Bulk import from a spreadsheet (CSV or Excel .xlsx): pupils with their parents and class, and
staff with their roles. The whole file is checked first and every problem is reported by its
spreadsheet row; nothing is saved until the file is clean, and then all of it is saved together.
Each record is made by the same services and checks as adding one by hand, so Student IDs,
parent accounts and the audit log are exactly the same."""

import csv
import io
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime

from django.db import transaction
from django.http import Http404, HttpResponse
from openpyxl import load_workbook
from rest_framework import serializers
from rest_framework.parsers import MultiPartParser
from rest_framework.response import Response

from apps.academics.models import SchoolClass
from apps.accounts.identifiers import normalise_identifier
from apps.accounts.models import User
from apps.people.models import Guardianship, Parent, Student
from apps.people.services import (
    STAFF_ROLES,
    add_emergency_contact,
    admit_student,
    create_parent,
    create_staff,
    enrol,
    link_guardian,
)
from apps.platform_admin.services import issue_password
from apps.school_admin.dashboard import current_year
from apps.school_admin.parents import ParentSerializer
from apps.school_admin.pupils import NewPupilSerializer
from apps.school_admin.staff import NewStaffSerializer
from apps.school_admin.views import NOTICE_GIVEN, SchoolView

MAX_BYTES = 2 * 1024 * 1024
# shortcut: one request must finish within the server's 60-second limit, so big lists are split
# into several files; move the saving to a background job if schools find that tiresome.
MAX_ROWS = 500
MAX_COLUMNS = 40
MAX_SCANNED = 50_000  # rows read before giving up on a sheet full of empty formatted rows

PARENT_PARTS = (
    ("first_name", "first name"),
    ("surname", "surname"),
    ("phone", "phone"),
    ("relationship", "relationship"),
    ("email", "email"),
)
PUPIL_COLUMNS = (
    ("first_name", "First name", True),
    ("middle_name", "Middle name", False),
    ("surname", "Surname", True),
    ("date_of_birth", "Date of birth", True),
    ("gender", "Gender", True),
    ("school_class", "Class", False),
    ("admission_date", "Admission date", False),
    ("previous_school", "Previous school", False),
    ("address", "Address", False),
    ("pickup_point", "Pick-up point", False),
    *(
        (f"parent{n}_{key}", f"Parent {n} {label}", False)
        for n in (1, 2)
        for key, label in PARENT_PARTS
    ),
)
STAFF_COLUMNS = (
    ("first_name", "First name", True),
    ("surname", "Surname", True),
    ("username", "Username", True),
    ("roles", "Roles", True),
    ("phone", "Phone", False),
    ("email", "Email", False),
    ("staff_number", "Staff number", False),
    ("qualifications", "Qualifications", False),
    ("licence_number", "Licence number", False),
    ("emergency_contact", "Emergency contact", False),
)
COLUMNS = {"pupils": PUPIL_COLUMNS, "staff": STAFF_COLUMNS}

UNREADABLE = (
    "We couldn't read this file. Save it as an Excel workbook (.xlsx) or as CSV, then try again."
)
OLD_EXCEL = (
    "This is an old Excel file (.xls). Save it as an Excel workbook (.xlsx) or as CSV, then try "
    "again."
)
DATE_FORMATS = ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%d/%m/%y")
GENDERS = {"m": "M", "male": "M", "boy": "M", "f": "F", "female": "F", "girl": "F"}
RELATIONSHIPS = {
    "mother": "MOTHER",
    "mum": "MOTHER",
    "mom": "MOTHER",
    "father": "FATHER",
    "dad": "FATHER",
    "guardian": "GUARDIAN",
    "other": "OTHER",
}


class SheetProblem(Exception):
    """Something wrong with the file as a whole, reported against its first row."""


def squash(text: str) -> str:
    """Compare headers, classes and roles the way people type them: "Crèche" == "creche "."""
    plain = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", plain.lower())


# ---- Reading the file ---------------------------------------------------------------------


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # Excel keeps 245550101 as a number, without its leading 0
    return str(value).strip()


def _keep(rows: list[list[str]], row: list[str]) -> bool:
    """Add a row; False once there are more filled rows than allowed (no need to read on)."""
    rows.append(row)
    filled = sum(1 for each in rows[1:] if any(each))
    return filled <= MAX_ROWS and len(rows) < MAX_SCANNED


def _excel_rows(data: bytes) -> list[list[str]]:
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # a damaged workbook, or another kind of zip file
        raise SheetProblem(UNREADABLE) from exc
    try:
        rows: list[list[str]] = []
        for row in book.worksheets[0].iter_rows(values_only=True, max_col=MAX_COLUMNS):
            if not _keep(rows, [_cell(value) for value in row]):
                break
        return rows
    finally:
        book.close()


def _csv_rows(data: bytes) -> list[list[str]]:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):  # Excel saves CSV in Windows letters
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    first_line = text.split("\n", 1)[0]
    delimiter = max((",", ";", "\t"), key=first_line.count)  # ";" where "," is a decimal mark
    rows: list[list[str]] = []
    try:
        for row in csv.reader(io.StringIO(text), delimiter=delimiter):
            if not _keep(rows, [cell.strip() for cell in row[:MAX_COLUMNS]]):
                break
    except csv.Error as exc:
        raise SheetProblem(UNREADABLE) from exc
    return rows


def read_rows(upload) -> list[list[str]]:
    data = upload.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise SheetProblem("The file is larger than 2 MB. Split the list into smaller files.")
    if data.startswith(b"PK"):  # .xlsx is a zip file
        return _excel_rows(data)
    if data.startswith(b"\xd0\xcf\x11\xe0"):
        raise SheetProblem(OLD_EXCEL)
    return _csv_rows(data)


def records(kind: str, rows: list[list[str]]) -> list[tuple[int, dict]]:
    """(spreadsheet row number, {column key: text}) for every filled row under the headers."""
    columns = COLUMNS[kind]
    if not rows:
        raise SheetProblem("The file is empty.")
    lookup = {}
    for key, header, _ in columns:
        lookup[squash(header)] = key
        lookup[squash(key)] = key
    positions: dict[str, int] = {}
    for index, title in enumerate(rows[0]):
        key = lookup.get(squash(title))
        if key and key not in positions:
            positions[key] = index
    missing = [header for key, header, required in columns if required and key not in positions]
    if missing:
        raise SheetProblem(
            f"Add {'this column' if len(missing) == 1 else 'these columns'} to the first row: "
            f"{', '.join(missing)}. Download the template to see every column."
        )
    found = []
    for number, row in enumerate(rows[1:], start=2):  # row 1 holds the headers, as in Excel
        values = {key: row[index] if index < len(row) else "" for key, index in positions.items()}
        if any(values.values()):
            found.append((number, values))
    if not found:
        raise SheetProblem("There are no rows to import under the headers.")
    if len(found) > MAX_ROWS:
        raise SheetProblem(
            f"Import at most {MAX_ROWS} rows at a time. Split the list into several files."
        )
    return found


# ---- Checking rows ------------------------------------------------------------------------


def _date(text: str) -> date:
    for pattern in DATE_FORMATS:
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    raise ValueError("write dates as DD/MM/YYYY, for example 14/03/2018.")


def _serializer_problems(serializer, headers: dict, reported: set) -> list[str]:
    """The first message for each column not already reported, named by its header."""
    if serializer.is_valid():
        return []
    return [
        f"{headers.get(name, name)}: {messages[0]}"
        for name, messages in serializer.errors.items()
        if name not in reported
    ]


def _missing(values: dict, columns, labels=None) -> tuple[list[str], set]:
    problems, reported = [], set()
    for key, header, required in columns:
        if required and not values.get(key):
            problems.append(f"{labels[key] if labels else header} is missing.")
            reported.add(key)
    return problems, reported


@dataclass
class PupilPlan:
    pupil: dict
    school_class: SchoolClass | None
    parents: list[tuple[str, dict, str]]  # (phone, parent fields, relationship), payer first


@dataclass
class Result:
    problems: list[dict]
    summary: dict
    plans: list
    existing: dict


PUPIL_HEADERS = {key: header for key, header, _ in PUPIL_COLUMNS}
PUPIL_FIELDS = (
    "first_name",
    "middle_name",
    "surname",
    "previous_school",
    "address",
    "pickup_point",
)


def _parent(number: int, values: dict) -> tuple[list[str], tuple | None]:
    prefix = f"parent{number}_"
    given = {key: values.get(prefix + key, "") for key, _ in PARENT_PARTS}
    if not any(given.values()):
        return [], None
    labels = {key: f"Parent {number} {label}" for key, label in PARENT_PARTS}
    required = [(key, labels[key], True) for key in ("first_name", "surname", "phone")]
    problems, reported = _missing(given, required, labels=labels)
    relationship = "GUARDIAN"
    if given["relationship"]:
        relationship = RELATIONSHIPS.get(squash(given["relationship"]), "")
        if not relationship:
            problems.append(f"{labels['relationship']}: write Mother, Father, Guardian or Other.")
    fields = {key: given[key] for key in ("first_name", "surname", "phone", "email") if given[key]}
    serializer = ParentSerializer(data=fields)
    problems += _serializer_problems(serializer, labels, reported)
    if problems:
        return problems, None
    data = dict(serializer.validated_data)
    return [], (data["phone"], data, relationship)


def check_pupils(school, found) -> Result:
    year = current_year(school)
    classes = {}
    if year:
        for school_class in SchoolClass.objects.for_school(school).filter(academic_year=year):
            classes[squash(school_class.name)] = school_class
    known = {
        (first.lower(), surname.lower(), born)
        for first, surname, born in Student.objects.for_school(school).values_list(
            "first_name", "surname", "date_of_birth"
        )
    }
    seen: dict[tuple, int] = {}
    problems, plans = [], []
    for number, values in found:
        row_problems, reported = _missing(values, PUPIL_COLUMNS)
        data = {key: values[key] for key in PUPIL_FIELDS if values.get(key)}
        for key in ("date_of_birth", "admission_date"):
            if values.get(key):
                try:
                    data[key] = _date(values[key]).isoformat()
                except ValueError as exc:
                    row_problems.append(f"{PUPIL_HEADERS[key]}: {exc}")
                    reported.add(key)
        if values.get("gender"):
            data["gender"] = GENDERS.get(squash(values["gender"]), "")
            if not data["gender"]:
                row_problems.append("Gender: write M or F.")
                reported.add("gender")
        school_class = None
        if values.get("school_class"):
            if year is None:
                row_problems.append(
                    "Class: set up the current school year first, or leave the class empty."
                )
            else:
                school_class = classes.get(squash(values["school_class"]))
                if school_class is None:
                    row_problems.append(
                        f"Class: No class called {values['school_class']} in {year.name}."
                    )
        serializer = NewPupilSerializer(data=data)
        row_problems += _serializer_problems(serializer, PUPIL_HEADERS, reported)
        if {"first_name", "surname", "date_of_birth"} <= data.keys() - reported:
            who = (
                data["first_name"].lower(),
                data["surname"].lower(),
                _date(data["date_of_birth"]),
            )
            if who in known:
                row_problems.append(
                    f"{data['first_name']} {data['surname']} is already a pupil at the school."
                )
            elif who in seen:
                row_problems.append(f"Same pupil as row {seen[who]}.")
            else:
                seen[who] = number
        parents = []
        for index in (1, 2):
            parent_problems, parent = _parent(index, values)
            row_problems += parent_problems
            if parent and any(parent[0] == other[0] for other in parents):
                row_problems.append("Parent 2 phone is the same as Parent 1's.")
            elif parent:
                parents.append(parent)
        problems += [{"row": number, "message": message} for message in row_problems]
        if not row_problems:
            plans.append(PupilPlan(dict(serializer.validated_data), school_class, parents))

    phones = {phone for plan in plans for phone, _, _ in plan.parents}
    existing = {
        parent.phone: parent
        for parent in Parent.objects.for_school(school).filter(phone__in=phones)
    }
    summary = {
        "pupils": len(plans),
        "new_parents": len(phones - existing.keys()),
        "existing_parents": len(phones & existing.keys()),
        "classes": sorted({plan.school_class.name for plan in plans if plan.school_class}),
        "unplaced": sum(1 for plan in plans if plan.school_class is None),
    }
    return Result(problems, summary, plans, existing)


def save_pupils(school, result: Result, by) -> dict:
    made: dict[str, Parent] = {}
    for plan in result.plans:
        pupil = admit_student(school, by=by, **plan.pupil)
        if plan.school_class:
            enrol(pupil, plan.school_class, start_date=pupil.admission_date, by=by)
        for index, (phone, fields, relationship) in enumerate(plan.parents):
            parent = result.existing.get(phone) or made.get(phone)
            if parent is None:
                parent = made[phone] = create_parent(school, by=by, **fields)
            first = index == 0  # the first parent listed pays and is the main contact
            link_guardian(
                parent, pupil, relationship, is_payer=first, is_primary_contact=first, by=by
            )
            # Parents are the people to call in an emergency, so a new pupil starts with them.
            add_emergency_contact(
                pupil,
                name=f"{parent.first_name} {parent.surname}",
                phone=parent.phone,
                relationship=Guardianship.Relationship(relationship).label,
            )
    return {}


STAFF_HEADERS = {key: header for key, header, _ in STAFF_COLUMNS}
ROLE_NAMES = {squash(name): name for name in STAFF_ROLES}


def check_staff(school, found) -> Result:
    seen: dict[str, int] = {}
    problems, plans, logins = [], [], {}
    for number, values in found:
        row_problems, reported = _missing(values, STAFF_COLUMNS)
        held, unknown = [], []
        for part in re.split(r"[;,/]+", values.get("roles", "")):
            if part.strip():
                name = ROLE_NAMES.get(squash(part))
                if name:
                    held.append(name)
                else:
                    unknown.append(part.strip())
        if unknown:
            row_problems.append(
                f"Roles: {', '.join(unknown)} {'is' if len(unknown) == 1 else 'are'} not a staff "
                f"role. Use {', '.join(sorted(STAFF_ROLES))}."
            )
            reported.add("roles")
        data = {key: values[key] for key in STAFF_HEADERS if key != "roles" and values.get(key)}
        data["roles"] = held
        serializer = NewStaffSerializer(data=data)
        row_problems += _serializer_problems(serializer, STAFF_HEADERS, reported)
        try:
            login = normalise_identifier(User.Kind.STAFF, values.get("username", ""))
        except ValueError:
            login = ""
        if login in seen:
            row_problems.append(f"Username: same username as row {seen[login]}.")
        elif login:
            seen[login] = number
            logins[number] = login
        problems += [{"row": number, "message": message} for message in row_problems]
        if not row_problems:
            plans.append((number, dict(serializer.validated_data)))

    taken = set(
        User.objects.filter(school=school, login__in=logins.values()).values_list(
            "login", flat=True
        )
    )
    for number, login in logins.items():
        if login in taken:
            problems.append(
                {"row": number, "message": f"Username: {login} is already used at this school."}
            )
    problems.sort(key=lambda problem: problem["row"])
    plans = [data for number, data in plans if logins.get(number) not in taken]
    return Result(problems, {"staff": len(plans)}, plans, {})


def save_staff(school, result: Result, by) -> dict:
    passwords = []
    for data in result.plans:
        data.pop("is_active", None)
        staff = create_staff(
            school, username=data.pop("username"), roles_held=data.pop("roles"), by=by, **data
        )
        passwords.append(
            {
                "name": f"{staff.first_name} {staff.surname}",
                "username": staff.user.login,
                "one_time_password": issue_password(staff.user),
            }
        )
    return {"passwords": passwords}


CHECKS = {"pupils": (check_pupils, save_pupils), "staff": (check_staff, save_staff)}


# ---- The screens' API ---------------------------------------------------------------------


class ImportView(SchoolView):
    parser_classes = [MultiPartParser]

    def post(self, request, kind):
        if kind not in CHECKS:
            raise Http404
        upload = request.FILES.get("file")
        if upload is None:
            raise serializers.ValidationError({"file": ["Choose a spreadsheet file."]})
        commit = request.data.get("commit") == "1"
        if commit and request.data.get("privacy_notice_given") != "1":
            raise serializers.ValidationError({"privacy_notice_given": [NOTICE_GIVEN]})
        body = {"kind": kind, "imported": False, "problems": [], "summary": None}
        try:
            found = records(kind, read_rows(upload))
        except SheetProblem as problem:
            body["problems"] = [{"row": 1, "message": str(problem)}]
            return Response(body)
        check, save = CHECKS[kind]
        result = check(self.school, found)
        body.update(problems=result.problems, summary=result.summary)
        if commit and not result.problems:
            with transaction.atomic():
                body.update(save(self.school, result, request.user), imported=True)
        return Response(body)


class ImportTemplateView(SchoolView):
    """The column headers as a CSV file that opens in Excel (the mark makes Excel read UTF-8)."""

    def get(self, request, kind):
        if kind not in COLUMNS:
            raise Http404
        headers = ",".join(header for _, header, _ in COLUMNS[kind])
        response = HttpResponse("﻿" + headers + "\r\n", content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = f'attachment; filename="{kind}-template.csv"'
        return response
